"""
Utility functions for report generation and export
"""
import csv
import io
from datetime import datetime
from decimal import Decimal

from django.http import HttpResponse
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False


def export_to_csv(data, filename, headers=None):
    """
    Export data to CSV format
    
    Args:
        data: List of dictionaries or list of lists
        filename: Output filename
        headers: Optional list of header names
    
    Returns:
        HttpResponse with CSV content
    """
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    writer = csv.writer(response)
    
    # Write headers
    if headers:
        writer.writerow(headers)
    elif data and isinstance(data[0], dict):
        writer.writerow(data[0].keys())
    
    # Write data
    for row in data:
        if isinstance(row, dict):
            writer.writerow(row.values())
        else:
            writer.writerow(row)
    
    return response


def export_to_excel(data, filename, sheet_name='Sheet1', headers=None, title=None):
    if not HAS_OPENPYXL:
        return export_to_csv(data, filename.replace('.xlsx', '.csv'), headers=headers)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    current_row = 1
    
    # Add title if provided
    if title:
        ws.merge_cells(f'A1:{chr(64 + len(headers or data[0].keys()))}1')
        title_cell = ws['A1']
        title_cell.value = title
        title_cell.font = Font(size=16, bold=True)
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        current_row = 3
    
    # Write headers
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    
    if headers:
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
    elif data and isinstance(data[0], dict):
        for col_num, header in enumerate(data[0].keys(), 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
    
    current_row += 1
    
    # Write data
    for row_data in data:
        if isinstance(row_data, dict):
            values = list(row_data.values())
        else:
            values = row_data
        
        for col_num, value in enumerate(values, 1):
            cell = ws.cell(row=current_row, column=col_num)
            
            # Handle different data types
            if isinstance(value, Decimal):
                cell.value = float(value)
                cell.number_format = '#,##0.00'
            elif isinstance(value, datetime):
                cell.value = value
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            else:
                cell.value = value
        
        current_row += 1
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    return response


def export_to_pdf(data, filename, headers=None, title=None, orientation=None):
    if not HAS_REPORTLAB:
        return export_to_csv(data, filename.replace('.pdf', '.csv'), headers=headers)
    
    buffer = io.BytesIO()

    # Determine table headers
    table_headers = []
    if headers:
        table_headers = headers
    elif data and isinstance(data[0], dict):
        table_headers = list(data[0].keys())

    num_cols = len(table_headers) if table_headers else 1

    # Auto-switch to landscape for tables with more than 5 columns
    if orientation is None:
        orientation = 'landscape' if num_cols > 5 else 'portrait'

    # Set page size & margins
    is_landscape = orientation == 'landscape'
    pagesize = (A4[1], A4[0]) if is_landscape else A4
    left_margin = 0.4 * inch
    right_margin = 0.4 * inch
    top_margin = 0.4 * inch
    bottom_margin = 0.4 * inch

    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=left_margin,
        rightMargin=right_margin,
        topMargin=top_margin,
        bottomMargin=bottom_margin
    )

    page_width, page_height = pagesize
    printable_width = page_width - (left_margin + right_margin)

    elements = []
    styles = getSampleStyleSheet()

    # Custom paragraph styles for cell content wrapping
    header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.whitesmoke,
        alignment=TA_CENTER
    )

    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor('#1e293b'),
        alignment=0
    )

    # Add title
    if title:
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#1e3a8a'),
            spaceAfter=14,
            alignment=TA_CENTER
        )
        elements.append(Paragraph(title, title_style))
        elements.append(Spacer(1, 0.15 * inch))

    # Prepare table data wrapped in Paragraphs
    table_data = []

    if table_headers:
        table_data.append([Paragraph(str(h), header_style) for h in table_headers])

    for row in data:
        row_cells = []
        if isinstance(row, dict):
            values = row.values()
        else:
            values = row

        for v in values:
            val_str = str(v) if v is not None else ''
            row_cells.append(Paragraph(val_str, cell_style))
        table_data.append(row_cells)

    # Proportional column width calculations to guarantee 100% fit within printable_width
    if num_cols > 0:
        weights = []
        for i in range(num_cols):
            h_name = str(table_headers[i]).lower() if i < len(table_headers) else ''
            if 'email' in h_name or 'name' in h_name or 'description' in h_name:
                weights.append(1.8)
            elif 'id' in h_name or 'date' in h_name or 'reference' in h_name:
                weights.append(1.2)
            else:
                weights.append(1.0)

        total_weight = sum(weights)
        col_widths = [(w / total_weight) * printable_width for w in weights]
    else:
        col_widths = None

    table = Table(table_data, colWidths=col_widths, repeatRows=1)

    table_style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 1), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ])

    table.setStyle(table_style)
    elements.append(table)

    # Footer
    elements.append(Spacer(1, 0.2 * inch))
    footer_style = ParagraphStyle(
        'Footer',
        parent=styles['Normal'],
        fontSize=8,
        textColor=colors.HexColor('#64748b'),
        alignment=TA_RIGHT
    )
    footer_text = f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    elements.append(Paragraph(footer_text, footer_style))

    doc.build(elements)

    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(pdf)

    return response


def format_currency(value):
    """Format decimal as currency"""
    if value is None:
        return '₦0.00'
    return f'₦{value:,.2f}'


def format_percentage(value):
    """Format decimal as percentage"""
    if value is None:
        return '0%'
    return f'{value:.1f}%'
